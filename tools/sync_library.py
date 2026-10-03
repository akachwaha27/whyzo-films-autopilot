"""Whyzo Films - PC sync.

Keeps a folder on your PC up to date with everything the bot makes:

  <your folder>/
    Videos/                     every published video, named "<date> - <title>.mp4"
    Thumbnails/                 the matching thumbnails
    Shorts Library.xlsx         sheets: Videos, Idea Lists, Comments, Daily Totals
    Dashboard.html              open in your browser to monitor everything
    _app/                       this script, its settings and a log

First time (once):   python sync_library.py --setup
  asks for the folder, installs what it needs, schedules a daily Windows task, runs the first sync.
Any time:            python sync_library.py          (sync now)
Remove the task:     python sync_library.py --uninstall
"""
import argparse
import datetime as dt
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time

REPO_DEFAULT = "akachwaha27/whyzo-films-autopilot"
TASK_NAME = "Whyzo Films Sync"
HERE = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(HERE, "sync_config.json")
DEPS = {"requests": "requests", "openpyxl": "openpyxl"}


# ------------------------------------------------------------------ setup helpers
def ensure_deps():
    missing = []
    for mod, pkg in DEPS.items():
        try:
            __import__(mod)
        except ImportError:
            missing.append(pkg)
    if missing:
        print("Installing:", ", ".join(missing))
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "--user", *missing])


def load_config():
    try:
        with open(CONFIG_FILE) as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def setup():
    default = load_config().get("folder") or os.path.join(os.path.expanduser("~"), "Documents", "Whyzo Films")
    folder = input(f"Folder for your videos, Excel and dashboard [{default}]: ").strip().strip('"') or default
    folder = os.path.abspath(os.path.expandvars(os.path.expanduser(folder)))
    repo = input(f"GitHub repo [{REPO_DEFAULT}]: ").strip() or REPO_DEFAULT
    at = input("Daily sync time, 24h HH:MM [21:30]: ").strip() or "21:30"
    if not re.fullmatch(r"\d{1,2}:\d{2}", at):
        raise SystemExit("Time must look like 21:30")
    ensure_deps()

    app = os.path.join(folder, "_app")
    os.makedirs(app, exist_ok=True)
    target = os.path.join(app, "sync_library.py")
    if os.path.abspath(__file__) != os.path.abspath(target):
        shutil.copy2(__file__, target)
    with open(os.path.join(app, "sync_config.json"), "w") as f:
        json.dump({"folder": folder, "repo": repo, "time": at}, f, indent=2)

    if os.name == "nt":
        register_task(target, at)
    else:
        print("Not on Windows: schedule this yourself, e.g. cron:", f"{sys.executable} {target}")
    print("\nRunning the first sync...")
    subprocess.call([sys.executable, target])
    dash = os.path.join(folder, "Dashboard.html")
    if os.path.exists(dash) and os.name == "nt":  # a browser, even if .html opens in Notepad
        if subprocess.call(f'start "" chrome "{dash}"', shell=True) and \
                subprocess.call(f'start "" msedge "{dash}"', shell=True):
            os.startfile(dash)  # noqa: S606
    print(f"\nDone. Everything lives in: {folder}")


def register_task(script, at):
    pyw = os.path.join(os.path.dirname(sys.executable), "pythonw.exe")
    exe = pyw if os.path.exists(pyw) else sys.executable
    ps = f"""
$a = New-ScheduledTaskAction -Execute '{exe}' -Argument '"{script}"' -WorkingDirectory '{os.path.dirname(script)}'
$t1 = New-ScheduledTaskTrigger -Daily -At '{at}'
$t2 = New-ScheduledTaskTrigger -AtLogOn -User $env:USERNAME
$s = New-ScheduledTaskSettingsSet -StartWhenAvailable -AllowStartIfOnBatteries -DontStopIfGoingOnBatteries -ExecutionTimeLimit (New-TimeSpan -Hours 2)
Register-ScheduledTask -TaskName '{TASK_NAME}' -Action $a -Trigger $t1,$t2 -Settings $s -Description 'Copies Whyzo Films videos, Excel and dashboard to your folder' -Force | Out-Null
"""
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
                       capture_output=True, text=True)
    if r.returncode != 0:
        print("Could not create the scheduled task:\n", r.stderr)
    else:
        print(f"Scheduled '{TASK_NAME}': daily at {at}, at sign-in, and as soon as possible if the PC was off.")


def uninstall():
    if os.name == "nt":
        subprocess.run(["schtasks", "/Delete", "/TN", TASK_NAME, "/F"])
    print("Scheduled task removed. Your folder and files are untouched.")


# ------------------------------------------------------------------ sync
def log(folder, msg):
    line = f"{dt.datetime.now():%Y-%m-%d %H:%M:%S}  {msg}"
    print(line)
    try:
        os.makedirs(os.path.join(folder, "_app"), exist_ok=True)
        with open(os.path.join(folder, "_app", "sync.log"), "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except OSError:
        pass


def safe_name(s, limit=90):
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", s or "untitled").strip().rstrip(".")
    s = re.sub(r"\s+", " ", s)
    return s[:limit].rstrip() or "untitled"


def fetch_json(session, repo, name):
    url = f"https://raw.githubusercontent.com/{repo}/main/archive/{name}.json"
    r = session.get(url, params={"t": int(time.time())}, timeout=60)
    if r.status_code == 404:
        return {}
    r.raise_for_status()
    return r.json()


def download(session, url, dest):
    tmp = dest + ".part"
    with session.get(url, stream=True, timeout=300) as r:
        if r.status_code == 404:
            return False
        r.raise_for_status()
        with open(tmp, "wb") as f:
            for chunk in r.iter_content(1 << 20):
                f.write(chunk)
    os.replace(tmp, dest)
    return True


def recent_runs(session, repo):
    try:
        r = session.get(f"https://api.github.com/repos/{repo}/actions/runs", params={"per_page": 30}, timeout=30)
        r.raise_for_status()
        return [{"name": x.get("display_title") or x.get("name"), "event": x.get("event"),
                 "status": x.get("status"), "conclusion": x.get("conclusion"),
                 "started": x.get("run_started_at") or x.get("created_at"), "url": x.get("html_url")}
                for x in r.json().get("workflow_runs", [])]
    except Exception:  # noqa: BLE001
        return []


def to_local(iso):
    if not iso:
        return None
    try:
        d = dt.datetime.fromisoformat(iso.replace("Z", "+00:00"))
        return d.astimezone().replace(tzinfo=None) if d.tzinfo else d
    except ValueError:
        return None


def run_sync():
    cfg = load_config()
    folder = cfg.get("folder")
    if not folder:
        raise SystemExit("Not set up yet. Run:  python sync_library.py --setup")
    repo = cfg.get("repo", REPO_DEFAULT)
    ensure_deps()
    import requests
    s = requests.Session()
    s.headers["User-Agent"] = "whyzo-films-sync"

    videos = fetch_json(s, repo, "videos")
    ideas = fetch_json(s, repo, "ideas")
    totals = fetch_json(s, repo, "stats_history")
    runs = recent_runs(s, repo)

    vdir, tdir = os.path.join(folder, "Videos"), os.path.join(folder, "Thumbnails")
    os.makedirs(vdir, exist_ok=True)
    os.makedirs(tdir, exist_ok=True)
    for v in videos.values():  # show the title as it is on YouTube now (you may have renamed it in Studio)
        yt_title = (v.get("youtube_stats") or {}).get("title")
        if yt_title and yt_title != v.get("title"):
            v["original_title"], v["title"] = v.get("title"), yt_title
    index_file = os.path.join(folder, "_app", "file_index.json")
    try:
        with open(index_file) as f:
            index = json.load(f)
    except (OSError, ValueError):
        index = {}
    got, used = 0, set()
    for vid, v in sorted(videos.items()):
        base = f"{v.get('date', vid[:10])} - {safe_name(v.get('title'))}"
        n = 2
        while base.lower() in used:  # same title twice on one day
            base = f"{v.get('date', vid[:10])} - {safe_name(v.get('title'))} ({n})"
            n += 1
        used.add(base.lower())
        v["_video_file"] = v["_thumb_file"] = ""
        for key, d, ext, field in (("file", vdir, ".mp4", "_video_file"), ("thumb", tdir, ".jpg", "_thumb_file")):
            ref = v.get(key) or {}
            dest = os.path.join(d, base + ext)
            olds = [index.get(f"{vid}{ext}")]
            if v.get("original_title"):
                olds.append(os.path.join(d, f"{v.get('date', vid[:10])} - {safe_name(v['original_title'])}{ext}"))
            for prev in olds:
                if prev and prev != dest and os.path.exists(prev) and not os.path.exists(dest):
                    os.replace(prev, dest)  # title changed in Studio: rename your copy to match
            if os.path.exists(dest):
                v[field] = dest
                index[f"{vid}{ext}"] = dest
                continue
            if v.get("status") != "published" or ref.get("type") != "release":
                continue
            url = f"https://github.com/{repo}/releases/download/{ref['tag']}/{ref['name']}"
            try:
                if download(s, url, dest):
                    v[field] = dest
                    index[f"{vid}{ext}"] = dest
                    got += key == "file"
                else:
                    v[field + "_missing"] = True
            except Exception as e:  # noqa: BLE001
                log(folder, f"download failed for {base}{ext}: {e}")
    os.makedirs(os.path.dirname(index_file), exist_ok=True)
    with open(index_file, "w") as f:
        json.dump(index, f, indent=1, ensure_ascii=False)
    xlsx = write_excel(folder, videos, ideas, totals)
    write_dashboard(folder, videos, ideas, totals, runs, repo)
    log(folder, f"sync ok: {len(videos)} videos in library, {got} new downloaded, "
                f"{sum(len(b.get('ideas', [])) for b in ideas.values())} ideas, Excel -> {os.path.basename(xlsx)}")


# ------------------------------------------------------------------ Excel
def _file_url(folder, path):
    root = os.getenv("SYNC_LINK_ROOT")  # only used when building the files on another machine
    if root:
        path = root.rstrip("\\/") + "\\" + os.path.relpath(path, folder).replace("/", "\\")
    return "file:///" + path.replace("\\", "/")


def write_excel(folder, videos, ideas, totals):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F2937")
    link_font = Font(color="0563C1", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    def sheet(ws, columns, rows):
        ws.append([c[0] for c in columns])
        for i, (name, width, kind) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = width
        for r in rows:
            ws.append([x if not isinstance(x, tuple) else x[0] for x in r])
            rn = ws.max_row
            for i, (x, (name, width, kind)) in enumerate(zip(r, columns), 1):
                c = ws.cell(row=rn, column=i)
                c.alignment = wrap if kind == "wrap" else top
                if kind == "date" and isinstance(x, dt.datetime):
                    c.number_format = "yyyy-mm-dd hh:mm"
                elif kind == "day":
                    c.number_format = "yyyy-mm-dd"
                elif kind == "int":
                    c.number_format = "#,##0"
                if isinstance(x, tuple) and x[1]:  # (text, link)
                    c.hyperlink = x[1]
                    c.font = link_font
        ws.freeze_panes = "B2" if len(columns) > 6 else "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30

    def plink(v, p):
        info = (v.get("platforms") or {}).get(p) or {}
        if info.get("link"):
            return ("Posted", info["link"])
        st = info.get("status") or ""
        return ("Not connected" if (not st or "skipped" in st) else st[:60])

    def day(s):
        try:
            return dt.datetime.strptime(str(s)[:10], "%Y-%m-%d")
        except (TypeError, ValueError):
            return s

    rows = []
    for vid, v in sorted(videos.items(), key=lambda kv: (kv[1].get("uploaded_at") or kv[0]), reverse=True):
        ys = v.get("youtube_stats") or {}
        yt = ((v.get("platforms") or {}).get("YouTube") or {})
        rows.append([
            day(v.get("date")), to_local(v.get("uploaded_at")), v.get("title", ""), v.get("format", ""),
            v.get("status", ""),
            ("Open on YouTube", yt["link"]) if yt.get("link") else (yt.get("status") or ""),
            ys.get("privacy", ""), ys.get("views"), ys.get("likes"), ys.get("comments"),
            plink(v, "Facebook"), plink(v, "Instagram"), plink(v, "TikTok"),
            v.get("description", ""), " ".join(v.get("hashtags") or []), ", ".join(v.get("tags") or []),
            v.get("pinned_comment", ""), v.get("script", ""), v.get("hook_text", ""), v.get("thumbnail_text", ""),
            v.get("primary_keyword", ""), v.get("voice", ""), v.get("angle", ""), v.get("trend_source", ""),
            v.get("virality_score"), v.get("approved_by", ""), to_local(v.get("scheduled_for")),
            "\n".join(v.get("credits") or []),
            (os.path.basename(v["_video_file"]), _file_url(folder, v["_video_file"])) if v.get("_video_file")
            else ("Not stored any more" if v.get("_video_file_missing") else ""),
            (os.path.basename(v["_thumb_file"]), _file_url(folder, v["_thumb_file"])) if v.get("_thumb_file") else "",
            vid,
        ])
    wb.active.title = "Videos"
    sheet(wb.active, [
        ("Date", 11, "day"), ("Uploaded (your time)", 16, "date"), ("Title", 40, "wrap"), ("Format", 11, ""),
        ("Status", 12, ""), ("YouTube", 16, ""), ("YouTube visibility", 12, ""), ("Views", 9, "int"),
        ("Likes", 8, "int"), ("Comments", 10, "int"), ("Facebook", 12, ""), ("Instagram", 12, ""),
        ("TikTok", 14, ""), ("Description", 60, "wrap"), ("Hashtags", 28, "wrap"), ("Tags", 50, "wrap"),
        ("Pinned comment", 40, "wrap"), ("Script", 70, "wrap"), ("Hook text", 28, "wrap"),
        ("Thumbnail text", 20, "wrap"), ("Main keyword", 22, ""), ("Voice", 24, ""), ("Idea angle", 50, "wrap"),
        ("Trend source", 22, ""), ("Virality score", 10, ""), ("Approved by", 11, ""),
        ("Scheduled public time", 18, "date"), ("Credits", 60, "wrap"), ("Video file", 40, ""),
        ("Thumbnail file", 30, ""), ("ID", 14, ""),
    ], rows)

    by_id = videos
    irows = []
    for d in sorted(ideas, reverse=True):
        b = ideas[d]
        for i in b.get("ideas", []):
            made = by_id.get(i.get("video_id") or "", {})
            irows.append([
                day(d), to_local(b.get("sent_at")), i.get("rank"), i.get("title", ""), i.get("format", ""),
                i.get("virality_score"), "Yes" if i.get("picked") else "No",
                made.get("status", "") if made else "", i.get("angle", ""), i.get("why", ""), i.get("trend_source", ""),
            ])
    sheet(wb.create_sheet("Idea Lists"), [
        ("Date", 11, "day"), ("Sent to Telegram", 16, "date"), ("#", 4, ""), ("Idea", 45, "wrap"),
        ("Format", 11, ""), ("Virality score", 10, ""), ("Picked", 8, ""), ("Video status", 12, ""),
        ("Angle", 60, "wrap"), ("Why it should work", 60, "wrap"), ("Trend source", 24, ""),
    ], irows)

    crows = []
    for vid, v in sorted(videos.items(), reverse=True):
        for c in v.get("latest_comments") or []:
            crows.append([day(v.get("date")), v.get("title", ""), c.get("author", ""), c.get("text", ""),
                          c.get("likes", 0), c.get("replies", 0), to_local(c.get("published"))])
    sheet(wb.create_sheet("Comments"), [
        ("Video date", 11, "day"), ("Video", 40, "wrap"), ("Author", 22, ""), ("Comment", 80, "wrap"),
        ("Likes", 8, "int"), ("Replies", 8, "int"), ("Posted (your time)", 16, "date"),
    ], crows)

    trows = [[day(d), t.get("videos"), t.get("views"), t.get("likes"), t.get("comments")]
             for d, t in sorted(totals.items(), reverse=True) if not d.startswith("_")]
    sheet(wb.create_sheet("Daily Totals"), [
        ("Date", 11, "day"), ("Videos on YouTube", 12, "int"), ("Total views", 12, "int"),
        ("Total likes", 12, "int"), ("Total comments", 12, "int"),
    ], trows)

    out = os.path.join(folder, "Shorts Library.xlsx")
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=folder)
    os.close(fd)
    wb.save(tmp)
    for _ in range(3):
        try:
            os.replace(tmp, out)
            return out
        except PermissionError:  # open in Excel
            time.sleep(5)
    alt = os.path.join(folder, "Shorts Library (latest - close Excel to update the main file).xlsx")
    try:
        os.replace(tmp, alt)
    except PermissionError:
        os.remove(tmp)
    return alt


# ------------------------------------------------------------------ dashboard
def write_dashboard(folder, videos, ideas, totals, runs, repo):
    vids = []
    for vid, v in videos.items():
        vids.append({
            "id": vid, "date": v.get("date"), "title": v.get("title"), "format": v.get("format"),
            "status": v.get("status"), "uploaded": v.get("uploaded_at"), "scheduled": v.get("scheduled_for"),
            "stats": v.get("youtube_stats") or {}, "platforms": v.get("platforms") or {},
            "thumb": ("Thumbnails/" + os.path.basename(v["_thumb_file"])) if v.get("_thumb_file") else "",
            "file": ("Videos/" + os.path.basename(v["_video_file"])) if v.get("_video_file") else "",
            "comments": len(v.get("latest_comments") or []),
        })
    data = {"videos": vids, "ideas": ideas, "totals": {k: v for k, v in totals.items() if not k.startswith("_")},
            "runs": runs, "repo": repo, "synced": dt.datetime.now().astimezone().isoformat(timespec="minutes"),
            "stats_checked": totals.get("_last_run")}
    blob = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    page = DASHBOARD.replace("__DATA__", blob).replace("__REPO__", html.escape(repo))
    tmp = os.path.join(folder, "Dashboard.html.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(page)
    os.replace(tmp, os.path.join(folder, "Dashboard.html"))


DASHBOARD = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Whyzo Films Dashboard</title>
<style>
:root{color-scheme:light;--page:#f9f9f7;--surface:#fcfcfb;--ink:#0b0b0b;--ink2:#52514e;--muted:#898781;--grid:#e1e0d9;
--axis:#c3c2b7;--border:rgba(11,11,11,.10);--s1:#2a78d6;--s2:#eb6834;--s3:#1baf7a;--s4:#eda100;--s5:#e87ba4;--s6:#008300;
--good:#0ca30c;--warn:#fab219;--crit:#d03b3b;--link:#1c5cab}
@media (prefers-color-scheme:dark){:root:where(:not([data-theme="light"])){color-scheme:dark;--page:#0d0d0d;--surface:#1a1a19;
--ink:#fff;--ink2:#c3c2b7;--grid:#2c2c2a;--axis:#383835;--border:rgba(255,255,255,.10);--s1:#3987e5;--s2:#d95926;--s3:#199e70;
--s4:#c98500;--s5:#d55181;--s6:#008300;--link:#86b6ef}}
:root[data-theme="dark"]{color-scheme:dark;--page:#0d0d0d;--surface:#1a1a19;--ink:#fff;--ink2:#c3c2b7;--grid:#2c2c2a;--axis:#383835;
--border:rgba(255,255,255,.10);--s1:#3987e5;--s2:#d95926;--s3:#199e70;--s4:#c98500;--s5:#d55181;--s6:#008300;--link:#86b6ef}
*{box-sizing:border-box}body{margin:0;background:var(--page);color:var(--ink);font:14px/1.45 system-ui,-apple-system,"Segoe UI",sans-serif}
.wrap{max-width:1280px;margin:0 auto;padding:20px 16px 48px}
header{display:flex;flex-wrap:wrap;align-items:baseline;gap:8px 16px;margin-bottom:14px}
h1{font-size:22px;margin:0}h2{font-size:15px;margin:0 0 2px}.sub{color:var(--ink2);font-size:13px}
.filters{display:flex;gap:6px;flex-wrap:wrap;margin:6px 0 16px}
.filters button{border:1px solid var(--border);background:var(--surface);color:var(--ink);border-radius:999px;padding:5px 12px;font:inherit;cursor:pointer}
.filters button[aria-pressed=true]{background:var(--ink);color:var(--surface);font-weight:600}
.tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-bottom:12px}
.card{background:var(--surface);border:1px solid var(--border);border-radius:12px;padding:14px 16px}
.tile .label{color:var(--ink2);font-size:13px}.tile .val{font-size:28px;font-weight:600;margin-top:2px}
.tile.hero .val{font-size:48px;line-height:1.05}.tile .note{color:var(--muted);font-size:12px}
.grid2{display:grid;grid-template-columns:repeat(auto-fit,minmax(380px,1fr));gap:12px;margin-bottom:12px}.grid2>*,.tiles>*{min-width:0}
@media (max-width:440px){.grid2{grid-template-columns:1fr}}
svg text{fill:var(--muted);font-size:11px;font-variant-numeric:tabular-nums}svg .lab{fill:var(--ink2);font-size:12px}
svg .val{fill:var(--ink);font-size:12px}
.legend{display:flex;flex-wrap:wrap;gap:4px 14px;margin:6px 0 4px;color:var(--ink2);font-size:12px}
.legend i{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:5px;vertical-align:-1px}
.empty{color:var(--muted);padding:24px 0;text-align:center}
#tip{position:fixed;pointer-events:none;background:var(--surface);border:1px solid var(--border);border-radius:8px;padding:6px 10px;
box-shadow:0 4px 14px rgba(0,0,0,.15);font-size:12px;display:none;z-index:9;max-width:280px}#tip b{font-size:14px;display:block}
table{width:100%;border-collapse:collapse}th,td{text-align:left;padding:8px 6px;border-bottom:1px solid var(--grid);vertical-align:top}
th{color:var(--ink2);font-weight:600;font-size:12px;white-space:nowrap}td.n{text-align:right;font-variant-numeric:tabular-nums}
.tw{overflow-x:auto}.thumb{width:40px;height:71px;object-fit:cover;border-radius:6px;background:var(--grid);display:block}
a{color:var(--link)}.pill{display:inline-flex;align-items:center;gap:4px;font-size:12px;color:var(--ink2);margin:0 8px 2px 0;white-space:nowrap}
.dot{width:8px;height:8px;border-radius:50%;display:inline-block}
.ideas li{margin:0 0 8px}.ideas .meta{color:var(--muted);font-size:12px}
.runs li{list-style:none;margin:0 0 6px}.runs{padding:0;margin:0}
.tag{font-size:11px;border:1px solid var(--border);border-radius:999px;padding:1px 7px;color:var(--ink2);margin-left:6px}
</style></head><body><div class="wrap">
<header><h1>Whyzo Films</h1><span class="sub" id="synced"></span></header>
<div class="filters" role="group" aria-label="Date range" id="filters"></div>
<div class="tiles" id="tiles"></div>
<div class="grid2">
 <div class="card"><h2>YouTube views by video</h2><div class="sub">Top 10 in this period</div><div id="c-views"></div></div>
 <div class="card"><h2>Videos uploaded per day</h2><div class="sub">By format</div><div id="c-days"></div></div>
</div>
<div class="grid2">
 <div class="card"><h2>Channel views over time</h2><div class="sub">Daily total across all uploaded videos</div><div id="c-total"></div></div>
 <div class="card"><h2>Latest ideas</h2><div class="sub" id="ideas-date"></div><ol class="ideas" id="ideas"></ol></div>
</div>
<div class="grid2">
 <div class="card" style="grid-column:1/-1"><h2>Videos</h2><div class="sub">Newest first. Click a title to open it on YouTube; "Play" opens your local copy.</div>
 <div class="tw"><table id="vt"><thead><tr><th></th><th>Title</th><th>Uploaded</th><th>Format</th><th>Platforms</th>
 <th style="text-align:right">Views</th><th style="text-align:right">Likes</th><th style="text-align:right">Comments</th><th>File</th></tr></thead><tbody></tbody></table></div></div>
</div>
<div class="card"><h2>Automation health</h2><div class="sub">Last runs on GitHub Actions</div><ul class="runs" id="runs"></ul></div>
</div><div id="tip"></div>
<script id="data" type="application/json">__DATA__</script>
<script>
const D=JSON.parse(document.getElementById('data').textContent);
const FORMATS=['ranking','story','funny','quiz','tips','explainer'];
const FC={ranking:'--s1',story:'--s2',funny:'--s3',quiz:'--s4',tips:'--s5',explainer:'--s6'};
const css=v=>getComputedStyle(document.documentElement).getPropertyValue(v).trim();
const el=(t,a={},...kids)=>{const e=document.createElement(t);for(const[k,v]of Object.entries(a)){if(k==='text')e.textContent=v;else if(k==='style')e.style.cssText=v;else e.setAttribute(k,v)}kids.forEach(k=>k&&e.append(k));return e};
const NS='http://www.w3.org/2000/svg';const sv=(t,a={})=>{const e=document.createElementNS(NS,t);for(const[k,v]of Object.entries(a))e.setAttribute(k,v);return e};
const fmtN=n=>n==null?'–':n>=1e6?(n/1e6).toFixed(n>=1e7?0:1)+'M':n>=1e4?(n/1e3).toFixed(0)+'K':n.toLocaleString();
const fmtT=s=>{if(!s)return'';const d=new Date(s);return isNaN(d)?s:d.toLocaleString([], {month:'short',day:'numeric',hour:'numeric',minute:'2-digit'})};
const tip=document.getElementById('tip');
function showTip(ev,val,label){tip.replaceChildren(el('b',{text:val}),el('span',{text:label}));tip.style.display='block';
 const x=Math.min(ev.clientX+14,innerWidth-tip.offsetWidth-8),y=ev.clientY+14;tip.style.left=x+'px';tip.style.top=y+'px'}
function hideTip(){tip.style.display='none'}
function hit(node,val,label){node.setAttribute('tabindex','0');node.addEventListener('pointermove',e=>showTip(e,val,label));
 node.addEventListener('pointerleave',hideTip);node.addEventListener('focus',()=>{const r=node.getBoundingClientRect();showTip({clientX:r.right,clientY:r.top},val,label)});node.addEventListener('blur',hideTip)}
document.getElementById('synced').textContent='Synced '+fmtT(D.synced)+(D.stats_checked?' · YouTube stats checked '+fmtT(D.stats_checked):'');

const RANGES=[['7 days',7],['30 days',30],['90 days',90],['All time',0]];let days=30;
const fbar=document.getElementById('filters');
RANGES.forEach(([l,n])=>{const b=el('button',{text:l,'aria-pressed':String(n===days)});b.onclick=()=>{days=n;[...fbar.children].forEach(c=>c.setAttribute('aria-pressed',String(c===b)));render()};fbar.append(b)});
const inRange=d=>{if(!days)return true;const c=new Date();c.setDate(c.getDate()-days+1);return d>=c.toISOString().slice(0,10)};

function render(){
 const V=D.videos.filter(v=>inRange(v.date||''));
 const pub=V.filter(v=>v.status==='published');
 const has=pub.some(v=>v.stats.views!=null);const sum=f=>has?pub.reduce((a,v)=>a+(v.stats[f]||0),0):null;
 const ideaDays=Object.values(D.ideas).filter(b=>inRange(b.date));
 const ideasN=ideaDays.reduce((a,b)=>a+b.ideas.length,0),picked=ideaDays.reduce((a,b)=>a+b.ideas.filter(i=>i.picked).length,0);
 const T=document.getElementById('tiles');T.replaceChildren();
 const tile=(label,val,note,hero)=>T.append(el('div',{class:'card tile'+(hero?' hero':'')},el('div',{class:'label',text:label}),el('div',{class:'val',text:val}),note?el('div',{class:'note',text:note}):null));
 tile('YouTube views',fmtN(sum('views')),!has?'Stats appear after the next check (every 2 hours)':pub.some(v=>v.stats.privacy==='private')?'Private videos get no public views until the API audit passes':'',true);
 tile('Videos published',String(pub.length),V.length-pub.length?`${V.length-pub.length} not published (skipped, failed or waiting)`:'');
 tile('Likes',fmtN(sum('likes')));tile('Comments',fmtN(sum('comments')));
 tile('Ideas sent',String(ideasN),ideasN?`${picked} picked (${Math.round(picked/ideasN*100)}%)`:'');
 barsViews(pub);columns(pub);line();ideasList();table(V);runs();
}
function barsViews(pub){const box=document.getElementById('c-views');box.replaceChildren();
 const rows=pub.filter(v=>v.stats.views!=null).sort((a,b)=>b.stats.views-a.stats.views).slice(0,10);
 if(!rows.length){box.append(el('div',{class:'empty',text:'No YouTube stats yet. They refresh every 2 hours.'}));return}
 const W=560,rowH=30,L=210,R=56,H=rows.length*rowH+8,max=Math.max(1,...rows.map(r=>r.stats.views));
 const s=sv('svg',{viewBox:`0 0 ${W} ${H}`,width:'100%',role:'img','aria-label':'YouTube views by video'});
 s.append(sv('line',{x1:L,x2:L,y1:0,y2:H,stroke:css('--axis'),'stroke-width':1}));
 rows.forEach((r,i)=>{const y=i*rowH+6,w=Math.max(2,(W-L-R)*r.stats.views/max),t=r.title.length>32?r.title.slice(0,31)+'…':r.title;
  const lab=sv('text',{x:L-8,y:y+15,'text-anchor':'end',class:'lab'});lab.textContent=t;s.append(lab);
  s.append(sv('path',{d:`M${L} ${y+4}h${w-4}a4 4 0 0 1 4 4v8a4 4 0 0 1 -4 4h${-(w-4)}z`,fill:css('--s1')}));
  const val=sv('text',{x:L+w+6,y:y+16,class:'val'});val.textContent=fmtN(r.stats.views);s.append(val);
  const h=sv('rect',{x:0,y:y,width:W,height:rowH-2,fill:'transparent'});hit(h,r.stats.views.toLocaleString()+' views',r.title);s.append(h)});
 box.append(s)}
function columns(pub){const box=document.getElementById('c-days');box.replaceChildren();
 const byDay={};pub.forEach(v=>{const d=v.date;byDay[d]=byDay[d]||{};byDay[d][v.format]=(byDay[d][v.format]||0)+1});
 const ds=Object.keys(byDay).sort();if(!ds.length){box.append(el('div',{class:'empty',text:'No uploads in this period.'}));return}
 const used=FORMATS.filter(f=>pub.some(v=>v.format===f));
 const lg=el('div',{class:'legend'});used.forEach(f=>{const i=el('i');i.style.background=css(FC[f]);lg.append(el('span',{},i,document.createTextNode(f)))});box.append(lg);
 const W=560,H=200,L=28,B=22,top=8,max=Math.max(1,...ds.map(d=>Object.values(byDay[d]).reduce((a,b)=>a+b,0)));
 const step=(W-L)/ds.length,bw=Math.min(24,step*0.6);
 const s=sv('svg',{viewBox:`0 0 ${W} ${H}`,width:'100%',role:'img','aria-label':'Videos uploaded per day by format'});
 const y=v=>H-B-(H-B-top)*v/max;
 for(let t=0;t<=max;t+=Math.max(1,Math.ceil(max/4))){s.append(sv('line',{x1:L,x2:W,y1:y(t),y2:y(t),stroke:css('--grid'),'stroke-width':1}));const tx=sv('text',{x:L-6,y:y(t)+4,'text-anchor':'end'});tx.textContent=t;s.append(tx)}
 ds.forEach((d,i)=>{const x=L+i*step+(step-bw)/2;let acc=0;const segs=used.filter(f=>byDay[d][f]);
  segs.forEach((f,j)=>{const n=byDay[d][f],y0=y(acc),y1=y(acc+n),last=j===segs.length-1,h=y0-y1-(last?0:2);
   const p=last?`M${x} ${y0}v${-(h-4)}a4 4 0 0 1 4 -4h${bw-8}a4 4 0 0 1 4 4v${h-4}z`:`M${x} ${y0-0}h${bw}v${-h}h${-bw}z`;
   const seg=sv('path',{d:p,fill:css(FC[f])});hit(seg,n+' '+f+(n>1?' videos':' video'),d);s.append(seg);acc+=n});
  if(ds.length<=14||i%Math.ceil(ds.length/10)===0){const tx=sv('text',{x:x+bw/2,y:H-6,'text-anchor':'middle'});tx.textContent=d.slice(5);s.append(tx)}});
 s.append(sv('line',{x1:L,x2:W,y1:H-B,y2:H-B,stroke:css('--axis'),'stroke-width':1}));box.append(s)}
function line(){const box=document.getElementById('c-total');box.replaceChildren();
 const pts=Object.entries(D.totals).filter(([d])=>inRange(d)).sort();
 if(pts.length<2){box.append(el('div',{class:'empty',text:pts.length?'One day of data so far; the line appears after the second day.':'No data yet.'}));return}
 const W=560,H=200,L=44,R=50,B=22,top=10,vals=pts.map(p=>p[1].views||0),max=Math.max(1,...vals);
 const x=i=>L+(W-L-R)*i/(pts.length-1),y=v=>H-B-(H-B-top)*v/max;
 const s=sv('svg',{viewBox:`0 0 ${W} ${H}`,width:'100%',role:'img','aria-label':'Channel views over time'});
 [0,max/2,max].forEach(t=>{s.append(sv('line',{x1:L,x2:W-R,y1:y(t),y2:y(t),stroke:css('--grid'),'stroke-width':1}));const tx=sv('text',{x:L-6,y:y(t)+4,'text-anchor':'end'});tx.textContent=fmtN(Math.round(t));s.append(tx)});
 s.append(sv('path',{d:vals.map((v,i)=>(i?'L':'M')+x(i)+' '+y(v)).join(''),fill:'none',stroke:css('--s1'),'stroke-width':2,'stroke-linejoin':'round','stroke-linecap':'round'}));
 const li=vals.length-1;s.append(sv('circle',{cx:x(li),cy:y(vals[li]),r:4,fill:css('--s1'),stroke:css('--surface'),'stroke-width':2}));
 const ev=sv('text',{x:x(li)+8,y:y(vals[li])+4,class:'val'});ev.textContent=fmtN(vals[li]);s.append(ev);
 [0,li].forEach(i=>{const tx=sv('text',{x:x(i),y:H-6,'text-anchor':i?'end':'start'});tx.textContent=pts[i][0].slice(5);s.append(tx)});
 const cross=sv('line',{y1:top,y2:H-B,stroke:css('--axis'),'stroke-width':1,visibility:'hidden'});s.append(cross);
 const ov=sv('rect',{x:L,y:0,width:W-L-R,height:H,fill:'transparent'});
 ov.addEventListener('pointermove',e=>{const r=s.getBoundingClientRect(),px=(e.clientX-r.left)*W/r.width,i=Math.max(0,Math.min(li,Math.round((px-L)/(W-L-R)*li)));
  cross.setAttribute('x1',x(i));cross.setAttribute('x2',x(i));cross.setAttribute('visibility','visible');showTip(e,(pts[i][1].views||0).toLocaleString()+' views',pts[i][0])});
 ov.addEventListener('pointerleave',()=>{cross.setAttribute('visibility','hidden');hideTip()});s.append(ov);box.append(s)}
function ideasList(){const days=Object.keys(D.ideas).sort();const box=document.getElementById('ideas');box.replaceChildren();
 if(!days.length){box.append(el('li',{class:'empty',text:'No idea lists yet.'}));return}
 const b=D.ideas[days[days.length-1]];document.getElementById('ideas-date').textContent='Sent '+fmtT(b.sent_at);
 b.ideas.forEach(i=>box.append(el('li',{},el('span',{text:i.title}),i.picked?el('span',{class:'tag',text:'picked'}):null,
  el('div',{class:'meta',text:[i.format,i.virality_score!=null?'score '+i.virality_score+'/10':'',i.why].filter(Boolean).join(' · ')}))))}
const PLAT=['YouTube','Facebook','Instagram','TikTok'];
function pstate(p,priv){if(!p)return['Not connected','--muted'];if(p.link&&(priv==='private'||/^private/.test(p.status||'')))return['Uploaded (private)','--warn'];if(p.link)return['Posted','--good'];const s=p.status||'';
 if(s.includes('skipped'))return['Not connected','--muted'];if(s.includes('draft'))return['Draft in app','--warn'];if(s.includes('FAIL'))return['Failed','--crit'];return[s.slice(0,24),'--warn']}
function table(V){const tb=document.querySelector('#vt tbody');tb.replaceChildren();
 const rows=[...V].sort((a,b)=>(b.uploaded||b.id).localeCompare(a.uploaded||a.id));
 if(!rows.length){tb.append(el('tr',{},el('td',{colspan:'9',class:'empty',text:'No videos in this period.'})));return}
 rows.forEach(v=>{const yt=(v.platforms.YouTube||{}).link;
  const img=v.thumb?el('img',{class:'thumb',src:v.thumb,alt:'',loading:'lazy'}):el('div',{class:'thumb'});
  const title=yt?el('a',{href:yt,target:'_blank',rel:'noopener',text:v.title}):el('span',{text:v.title});
  const pl=el('div');PLAT.forEach(n=>{const[l,c]=pstate(v.platforms[n],n==='YouTube'?v.stats.privacy:null);const d=el('span',{class:'dot'});d.style.background=css(c);pl.append(el('span',{class:'pill',title:(v.platforms[n]||{}).status||''},d,document.createTextNode(n+': '+l)))});
  const st=v.status!=='published'?el('div',{class:'sub',text:'Status: '+v.status}):null;
  tb.append(el('tr',{},el('td',{},img),el('td',{},title,st),el('td',{text:fmtT(v.uploaded)||v.date}),el('td',{text:v.format||''}),el('td',{},pl),
   el('td',{class:'n',text:fmtN(v.stats.views)}),el('td',{class:'n',text:fmtN(v.stats.likes)}),el('td',{class:'n',text:fmtN(v.stats.comments)}),
   el('td',{},v.file?el('a',{href:v.file,text:'Play'}):el('span',{class:'sub',text:'—'}))))})}
function runs(){const box=document.getElementById('runs');box.replaceChildren();
 if(!D.runs.length){box.append(el('li',{class:'sub',text:'Could not reach GitHub during the last sync.'}));return}
 const last=D.runs.slice(0,8),done=D.runs.filter(r=>r.status==='completed'),ok=done.filter(r=>r.conclusion==='success').length;
 box.append(el('li',{class:'sub',text:`${ok} of ${done.length} recent finished runs succeeded.`}));
 last.forEach(r=>{const[l,c]=r.conclusion==='success'?['Succeeded','--good']:r.conclusion==='failure'?['Failed','--crit']:r.status!=='completed'?['Running','--warn']:[r.conclusion||'–','--muted'];
  const d=el('span',{class:'dot'});d.style.background=css(c);box.append(el('li',{},d,document.createTextNode(' '+l+' · '+fmtT(r.started)+' · '+(r.event||'')+' '),el('a',{href:r.url,target:'_blank',rel:'noopener',text:'details'})))})}
render();
matchMedia('(prefers-color-scheme: dark)').addEventListener('change',render);
</script></body></html>
"""


def main():
    ap = argparse.ArgumentParser(description="Sync Whyzo Films videos, Excel and dashboard to your PC")
    ap.add_argument("--setup", action="store_true", help="choose the folder and schedule the daily sync")
    ap.add_argument("--uninstall", action="store_true", help="remove the scheduled task")
    a = ap.parse_args()
    if a.setup:
        return setup()
    if a.uninstall:
        return uninstall()
    cfg = load_config()
    try:
        run_sync()
    except Exception as e:  # noqa: BLE001
        if cfg.get("folder"):
            log(cfg["folder"], f"sync FAILED: {e}")
        raise


if __name__ == "__main__":
    main()
